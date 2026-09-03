from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def pdf_oku(dosya):
    reader = PdfReader(dosya)
    parcalar = []
    for i, sayfa in enumerate(reader.pages, start=1):
        icerik = sayfa.extract_text()
        if icerik and icerik.strip():
            parcalar.append(f"[Sayfa {i}]\n{icerik.strip()}")
    return "\n\n".join(parcalar)


def docx_oku(dosya):
    belge = Document(dosya)
    govde = belge.element.body
    tablo_sirasi = iter(belge.tables)
    parcalar = []

    for eleman in govde.iterchildren():
        etiket = eleman.tag.split("}")[-1]

        if etiket == "p":
            yazi = "".join(
                d.text or "" for d in eleman.iter() if d.tag.endswith("}t")
            )
            if yazi.strip():
                parcalar.append(yazi.strip())

        elif etiket == "tbl":
            tablo = next(tablo_sirasi, None)
            if tablo is None:
                continue

            satirlar = list(tablo.rows)
            if not satirlar:
                continue

            basliklar = [h.text.strip() for h in satirlar[0].cells]
            metinler = [" | ".join(basliklar)]

            for satir in satirlar[1:]:
                hucreler = [h.text.strip() for h in satir.cells]
                cift = [
                    f"{basliklar[i]}: {hucreler[i]}"
                    for i in range(min(len(basliklar), len(hucreler)))
                    if hucreler[i]
                ]
                if cift:
                    metinler.append(" | ".join(cift))

            parcalar.append("\n".join(metinler))

    return "\n".join(parcalar)


def xlsx_oku(dosya):
    kitap = load_workbook(dosya, data_only=True)
    parcalar = []

    for sayfa in kitap.worksheets:
        satirlar = []
        for satir in sayfa.iter_rows(values_only=True):
            hucreler = [
                str(h).strip() for h in satir if h is not None and str(h).strip()
            ]
            if hucreler:
                satirlar.append(hucreler)

        if not satirlar:
            continue

        parcalar.append(f"[Sayfa: {sayfa.title}]")

        basliklar = None
        for hucreler in satirlar:
            if basliklar is None and len(hucreler) >= 2:
                basliklar = hucreler
                parcalar.append(" | ".join(basliklar))
                continue

            if basliklar and len(hucreler) >= 2:
                cift = [
                    f"{basliklar[i]}: {hucreler[i]}"
                    for i in range(min(len(basliklar), len(hucreler)))
                ]
                parcalar.append(" | ".join(cift))
            else:
                parcalar.append(" ".join(hucreler))

    return "\n".join(parcalar)


def pptx_oku(dosya):
    sunum = Presentation(dosya)
    parcalar = []

    for i, slayt in enumerate(sunum.slides, start=1):
        satirlar = [f"[Slayt {i}]"]

        for sekil in slayt.shapes:
            if sekil.has_text_frame:
                yazi = sekil.text_frame.text.strip()
                if yazi:
                    satirlar.append(yazi)

            if sekil.has_table:
                tablo = sekil.table
                for satir in tablo.rows:
                    hucreler = [h.text.strip() for h in satir.cells]
                    satirlar.append(" | ".join(hucreler))

        if slayt.has_notes_slide:
            not_metni = slayt.notes_slide.notes_text_frame.text.strip()
            if not_metni:
                satirlar.append(f"[Konusmaci notu] {not_metni}")

        if len(satirlar) > 1:
            parcalar.append("\n".join(satirlar))

    return "\n\n".join(parcalar)


def metin_oku(dosya):
    ham = dosya.read()
    if isinstance(ham, bytes):
        for kodlama in ("utf-8", "utf-8-sig", "cp1254", "latin-1"):
            try:
                return ham.decode(kodlama)
            except UnicodeDecodeError:
                continue
        return ham.decode("utf-8", errors="replace")
    return ham


OKUYUCULAR = {
    ".pdf": pdf_oku,
    ".docx": docx_oku,
    ".xlsx": xlsx_oku,
    ".pptx": pptx_oku,
    ".txt": metin_oku,
    ".md": metin_oku,
}